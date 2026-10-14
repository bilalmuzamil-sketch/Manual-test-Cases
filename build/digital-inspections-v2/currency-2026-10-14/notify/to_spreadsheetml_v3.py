#!/usr/bin/env python3
"""Convert the notification xlsx to Excel 2003 XML (SpreadsheetML): plain text, several tabs, uploadable as textContent."""
import re, sys, os
from xml.sax.saxutils import escape
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
H = os.path.dirname(os.path.abspath(__file__))
wb = load_workbook(f"{H}/DVI-V2-test-cases-with-changed-behaviour.xlsx")
A1 = re.compile(r"(?:'([^']+)'!)?\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d*))?")
def r1c1(f):
    # =COUNTIF('Tab'!M2:M,"x") -> =COUNTIF('Tab'!R2C13:R2000C13,"x"); COUNTA('Tab'!A2:A)
    def sub(m):
        tab, c1, r1, c2, r2 = m.groups(); pre = f"'{tab}'!" if tab else ""
        s = f"R{r1}C{column_index_from_string(c1)}"
        if c2: s += f":R{r2 or 2000}C{column_index_from_string(c2)}"
        return pre + s
    return re.sub(r"(?:'[^']+'!)\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d*)?", lambda m: sub(A1.fullmatch(m.group(0))), f)
HL = re.compile(r'=HYPERLINK\("([^"]+)","([^"]*)"\)')
out = ['<?xml version="1.0"?>', '<?mso-application progid="Excel.Sheet"?>',
       '<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet" xmlns:o="urn:schemas-microsoft-com:office:office" xmlns:x="urn:schemas-microsoft-com:office:excel" xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet">',
       '<Styles>',
       '<Style ss:ID="Default" ss:Name="Normal"><Alignment ss:Vertical="Top" ss:WrapText="1"/><Font ss:FontName="Arial" ss:Size="10"/></Style>',
       '<Style ss:ID="h"><Alignment ss:Vertical="Top" ss:WrapText="1"/><Font ss:FontName="Arial" ss:Size="10" ss:Bold="1" ss:Color="#FFFFFF"/><Interior ss:Color="#1F3864" ss:Pattern="Solid"/></Style>',
       '<Style ss:ID="t"><Alignment ss:Vertical="Top" ss:WrapText="1"/><Font ss:FontName="Arial" ss:Size="14" ss:Bold="1"/></Style>',
       '<Style ss:ID="b"><Alignment ss:Vertical="Top" ss:WrapText="1"/><Font ss:FontName="Arial" ss:Size="10" ss:Bold="1"/></Style>',
       '<Style ss:ID="l"><Alignment ss:Vertical="Top" ss:WrapText="1"/><Font ss:FontName="Arial" ss:Size="10" ss:Color="#1155CC" ss:Underline="Single"/></Style>',
       '</Styles>']
for ws in wb.worksheets:
    out.append(f'<Worksheet ss:Name="{escape(ws.title)}"><Table>')
    for col in range(1, ws.max_column + 1):
        letter = ws.cell(row=1, column=col).column_letter
        w = ws.column_dimensions[letter].width or 10
        out.append(f'<Column ss:Width="{round(w * 6)}"/>')
    readme = ws.title == "Read me"
    for r in range(1, ws.max_row + 1):
        cells = []; skip = False
        for c in range(1, ws.max_column + 1):
            idx = f' ss:Index="{c}"' if skip else ""; skip_prev = skip
            v = ws.cell(row=r, column=c).value
            if v is None or v == "": skip = True; continue
            v = str(v); m = HL.fullmatch(v)
            if m:
                cells.append(f'<Cell{idx} ss:StyleID="l" ss:HRef="{escape(m.group(1), {chr(34): "&quot;"})}"><Data ss:Type="String">{escape(m.group(2))}</Data></Cell>'); continue
            st = ' ss:StyleID="h"' if (r == 1 and not readme) else (' ss:StyleID="t"' if (readme and r == 1) else (' ss:StyleID="b"' if (readme and c == 1) else ""))
            if v.startswith("="):
                cells.append(f'<Cell{idx}{st} ss:Formula="{escape(r1c1(v), {chr(34): "&quot;"})}"><Data ss:Type="Number">0</Data></Cell>')
            else:
                cells.append(f'<Cell{idx}{st}><Data ss:Type="String">{escape(v)}</Data></Cell>')
        out.append("<Row>" + "".join(cells) + "</Row>")
    out.append("</Table>")
    if not readme:
        out.append('<WorksheetOptions xmlns="urn:schemas-microsoft-com:office:excel"><FreezePanes/><FrozenNoSplit/><SplitHorizontal>1</SplitHorizontal><TopRowBottomPane>1</TopRowBottomPane><ActivePane>2</ActivePane></WorksheetOptions>')
    out.append("</Worksheet>")
out.append("</Workbook>")
p = f"{H}/DVI-V2-test-cases-with-changed-behaviour.xml"; open(p, "w").write("\n".join(out)); print(os.path.getsize(p))
