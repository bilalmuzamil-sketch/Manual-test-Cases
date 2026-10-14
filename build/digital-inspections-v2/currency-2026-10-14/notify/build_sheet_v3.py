#!/usr/bin/env python3
"""DVI V2: one row per test case whose behaviour changed — before/after expected and steps, with spec dates."""
import json, os, datetime, re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
H = os.path.dirname(os.path.abspath(__file__))
B = json.load(open(f"{H}/behaviour-cases.json"))
PRD = "https://shopview.atlassian.net/wiki/pages/viewpage.action?pageId=768507905"
SLACK = "https://shopview.slack.com/archives/C0C39T5MLMB/p1791807912692899"
TR = lambda c: f"https://shopview.testrail.io/index.php?/cases/view/{c}"
link = lambda u, t: f'=HYPERLINK("{u}","{t}")'
F = "Arial"; HDR = PatternFill("solid", fgColor="1F3864"); thin = Side(style="thin", color="BFBFBF")
def nice(t):
    def one(v):
        try: return datetime.date.fromisoformat(v).strftime("%-d %b %Y")
        except Exception: return v
    return re.sub(r"\d{4}-\d{2}-\d{2}", lambda m: one(m.group(0)), t or "")
def lines(x): return "\n".join(x) if isinstance(x, list) else (x or "")
wb = Workbook(); rd = wb.active; rd.title = "Read me"
ws = wb.create_sheet("Test cases to update")
cols = ["Case", "Case title", "Automated?", "What changed in behaviour", "Spec rule and when the spec changed it",
        "Expected behaviour before", "Expected behaviour now", "Steps before", "Steps now", "Source"]
widths = [11, 32, 11, 38, 34, 50, 50, 50, 50, 12]
ws.append(cols)
for i, w in enumerate(widths, 1):
    c = ws.cell(row=1, column=i); c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = HDR
    c.alignment = Alignment(wrap_text=True, vertical="top"); ws.column_dimensions[c.column_letter].width = w
ws.freeze_panes = "A2"
for r in sorted(B["cases"], key=lambda r: r["case"]):
    rules = "\n".join(f"{x['code']}: {nice(x.get('date',''))}" + (f" ({x['entry']})" if x.get("entry") else "") for x in r["rules"])
    src = link(f"https://shopview.atlassian.net/browse/{r['rules'][0]['code']}", "Jira") if r["rules"] and r["rules"][0]["code"].startswith("SV-") else link(PRD, "Spec page")
    ws.append([link(TR(r["case"]), f"C{r['case']}"), r["case_title"], "Yes" if r.get("automated") else "No", r["why"], rules,
               lines(r["expected_before"]), lines(r["expected_now"]), lines(r["steps_before"]), lines(r["steps_now"]), src])
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name=F, color="1155CC", underline="single") if str(c.value).startswith("=HYPERLINK") else Font(name=F)
        c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = Border(bottom=thin)
# ---- New test cases needed
NT = json.load(open(f"{H}/new-tests.json"))["tests"] if os.path.exists(f"{H}/new-tests.json") else []
nt = wb.create_sheet("New test cases needed")
ncols = ["Proposed new test case", "Why it is needed (which change)", "Spec rule and when the spec changed it", "Preconditions", "Steps", "Expected behaviour", "Next to existing case", "Confirm with product first?"]
nw = [32, 42, 34, 40, 46, 46, 13, 30]
nt.append(ncols)
for i, w in enumerate(nw, 1):
    c = nt.cell(row=1, column=i); c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = HDR
    c.alignment = Alignment(wrap_text=True, vertical="top"); nt.column_dimensions[c.column_letter].width = w
nt.freeze_panes = "A2"
for t in NT:
    rules = "\n".join(f"{x['code']}: {nice(x.get('date',''))}" + (f" ({x['entry']})" if x.get("entry") else "") for x in t["rules"])
    rc = t.get("related_case")
    nt.append([t["title"], t["why_needed"], rules, "\n".join(t["preconditions"]), "\n".join(f"{i}. {x}" for i, x in enumerate(t["steps"], 1)),
               "\n".join(t["expected"]), link(TR(rc), f"C{rc}") if rc else "", t.get("needs_confirmation") or "No"])
for row in nt.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name=F, color="1155CC", underline="single") if str(c.value).startswith("=HYPERLINK") else Font(name=F)
        c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = Border(bottom=thin)
rd.column_dimensions["A"].width = 30; rd.column_dimensions["B"].width = 100
L = [("DVI V2 test cases whose expected behaviour has changed", ""), ("", ""),
     ("Hi Viktoria and Vladimir,", ""),
     ("", "We compared every Digital Inspection V2 test case with the current spec (version 43, 13 October 2026) and the later decisions in Jira. This sheet lists only the test cases where the expected behaviour has changed. Wording-only changes are left out."),
     ("", "For each test case it shows the expected behaviour and steps as they are in the case today (before), and as they should be now, with the spec rule and the date the spec changed it."),
     ("Viktoria", "The DVI V2 team sync of 12 October says \"QA on the feature squad owns writing and updating test cases\", so these updates are yours to make. We have not edited any test case."),
     ("Vladimir", "Test cases marked Automated = Yes may have automation that still checks the old behaviour."),
     ("", ""),
     ("Test cases to update", "=COUNTA('Test cases to update'!A2:A)"),
     ("Of which automated", "=COUNTIF('Test cases to update'!C2:C,\"Yes\")"),
     ("New test cases needed", "=COUNTA('New test cases needed'!A2:A)"),
     ("", "The tab \"New test cases needed\" lists behaviour that changed or was added and that no existing test case checks, with the change that makes each one necessary."),
     ("", ""),
     ("The spec", link(PRD, "Confluence page 768507905, Digital Inspections V2, version 43 (13 Oct 2026)")),
     ("When each rule changed", "From the Change Log at the end of the spec. Where a rule is not named in the change log, a date range is given."),
     ("Team sync, 12 Oct 2026", link(SLACK, "Milos Vasic in #fs-dvi-v2")),
     ("The test cases", link("https://shopview.testrail.io/index.php?/suites/view/1&group_id=6658", "TestRail folder: Digital Inspection V2")),
     ("", ""), ("Thank you,", "Bilal Muzamil")]
for a, b in L: rd.append([a, b])
for row in rd.iter_rows():
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(name=F, color="1155CC", underline="single") if str(c.value).startswith("=HYPERLINK") else Font(name=F, bold=(c.column == 1 and bool(c.value)))
rd["A1"].font = Font(name=F, bold=True, size=14)
out = f"{H}/DVI-V2-test-cases-with-changed-behaviour.xlsx"; wb.save(out); print(len(B["cases"]), "cases ->", out)
