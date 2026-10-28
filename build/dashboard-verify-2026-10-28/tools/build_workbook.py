"""Build the consolidated findings workbook (Summary tab + one tab per status + PO questions).

Usage: python3 build_workbook.py  → writes ../ShopView-Dashboard-and-Reports-Full-Check-Findings-2026-10-28.xlsx
"""
import os, re, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
from findings_data import FINDINGS, CHECKS  # noqa: E402

GH = 'https://github.com/bilalmuzamil-sketch/Manual-test-Cases/blob/claude/credit-memo-document-location-tov2x1/build/dashboard-verify-2026-10-28/'
TABS = [('READY', 'Ready to file', 'Passed every admissibility check. Waits only for your go-ahead, one ticket at a time.'),
        ('HELD', 'Held for PO', 'Draft written. The written source is silent or can be read two ways, so the PO answers first.'),
        ('KNOWN', 'Already reported', 'Already in Jira and still happening on this build. Not filed again.'),
        ('QA-LEAD', 'QA lead decision', 'Raised before, closed as Obsolete, still on screen. You decide whether to reopen.'),
        ('NOT-TICKET', 'Not a ticket', 'Only reachable through the API (screen blocks it, Rule 24), or rounding / observation.')]
HEAD = Font(bold=True, color='FFFFFF')
FILL = PatternFill('solid', fgColor='1F4E79')
WRAP = Alignment(wrap_text=True, vertical='top')


def sheet(ws, header, rows, widths):
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HEAD, FILL, WRAP
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    ws.freeze_panes = 'A2'


DOC = 'https://docs.google.com/document/d/1mEcB4cpsOqZXeim0vzmlbam2xSB02YshnyYRY0dk9s8/edit'


def po_rows():
    """Rows of the Chris Ward question doc (doc number, our id, title) — the doc is the PO sheet now."""
    f = os.path.join(ROOT, 'questions-for-chris-2026-10-28', 'QUESTIONS-MAP.md')
    out = []
    for line in open(f, encoding='utf-8'):
        c = [x.strip() for x in line.strip().strip('|').split('|')]
        if len(c) == 3 and c[0].isdigit():
            out.append([int(c[0]), c[1], c[2], DOC, ''])
    return out


def left_out_rows():
    f = os.path.join(ROOT, 'questions-for-chris-2026-10-28', 'LEFT-OUT-NOT-ON-SCREEN.md')
    out = []
    for line in open(f, encoding='utf-8'):
        c = [x.strip() for x in line.strip().strip('|').split('|')]
        if len(c) == 3 and c[0] not in ('Item', '---'):
            out.append(c)
    return out


def main():
    wb = Workbook()
    s = wb.active
    s.title = 'Summary'
    s.append(['ShopView — Dashboard and Reports full check, 28 Oct 2026'])
    s['A1'].font = Font(bold=True, size=14)
    s.append(['Build checked: staging v26.43.2-f8698cb · Role: Admin · Locations: Staging Heavy Duty - 9919, Staging Lethbridge - 4310, QB Location, DVI Test'])
    s.append(['TestRail: not touched (your instruction). Jira: nothing created (needs your go-ahead per ticket).'])
    s.append([])
    s.append(['Findings by outcome', 'How many', 'What it means'])
    for c in s[5]:
        c.font, c.fill = HEAD, FILL
    for key, name, meaning in TABS:
        s.append([name, sum(1 for f in FINDINGS if f['status'] == key), meaning])
    s.append([])
    s.append(['Checks run', 'Checks', 'Passed', 'Failed / held', 'Notes'])
    for c in s[s.max_row]:
        c.font, c.fill = HEAD, FILL
    for row in CHECKS:
        s.append(row)
    s.column_dimensions['A'].width = 48
    s.column_dimensions['B'].width = 12
    s.column_dimensions['C'].width = 70
    s.column_dimensions['D'].width = 14
    s.column_dimensions['E'].width = 60
    for row in s.iter_rows(min_row=5):
        for c in row:
            c.alignment = WRAP

    for key, name, _ in TABS:
        rows = [[f['id'], f['area'], f['report'], f['finding'], f['evidence'], f['source'], f['question'],
                 f['files'], GH + f['files']] for f in FINDINGS if f['status'] == key]
        sheet(wb.create_sheet(name), ['#', 'Area', 'Report / tile', 'What is wrong', 'Evidence', 'Written source',
                                       'PO question', 'Draft / file', 'Link'], rows, [9, 12, 26, 48, 60, 36, 11, 40, 40])
    sheet(wb.create_sheet('Questions for Chris'), ['Doc #', 'Our id', 'Question (report and point)', 'Google Doc', 'Answer'],
          po_rows(), [7, 9, 70, 50, 30])
    sheet(wb.create_sheet('Left out - not on screen'), ['Item', 'What it was', 'Why a customer cannot reach it'],
          left_out_rows(), [18, 60, 70])
    out = os.path.join(ROOT, 'ShopView-Dashboard-and-Reports-Full-Check-Findings-2026-10-28.xlsx')
    wb.save(out)
    print(out, {k: sum(1 for f in FINDINGS if f['status'] == k) for k, _, _ in TABS}, 'questions', len(po_rows()), 'left out', len(left_out_rows()))


if __name__ == '__main__':
    main()
